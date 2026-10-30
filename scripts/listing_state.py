"""Read-only XLSX title history and atomic, explicitly verified delivery state.

No workbook authoring, network APIs or product publication. Historical coverage
is limited to explicitly supplied workbooks and the selected local JSON ledger.
Keep state outside the skill source/installation tree, in a batch or ledger dir.
"""
import argparse
from contextlib import contextmanager
from copy import deepcopy
import hashlib
import io
import json
import os
from pathlib import Path
import re
import sys
import tempfile
import unicodedata
import uuid


class ListingStateError(ValueError):
    """An unsafe, ambiguous or stale listing state; no state was committed."""


def normalize_title(title):
    """Canonical historical comparison: NFKC, without invisible Cf characters."""
    if not isinstance(title,str) or not title.strip():
        raise ListingStateError('Title must be nonempty text')
    result=''.join(c for c in unicodedata.normalize('NFKC',title) if unicodedata.category(c)!='Cf').strip()
    if not result:raise ListingStateError('Title must contain visible text')
    return result


def _new_title(title):
    normalized=normalize_title(title)
    if any(unicodedata.category(c) in ('Cf','Cc') for c in title):
        raise ListingStateError('New titles cannot contain invisible format/control characters')
    return normalized


def _uid(value):
    if not isinstance(value,str) or not re.fullmatch(r'S[0-9]+',value):
        raise ListingStateError('Invalid or unconfirmed store UID')
    return value


def _batch(value):
    if (not isinstance(value,str) or not value or value!=value.strip() or len(value)>200
            or any(ord(c)<32 for c in value)):
        raise ListingStateError('A nonempty, explicit batch_id is required')
    return value


def _hash(value):
    if not isinstance(value,str) or not re.fullmatch('[0-9a-f]{64}',value):
        raise ListingStateError('Expected a SHA-256 digest')
    return value


def _link_ids(values,count):
    if (not isinstance(values,(list,tuple)) or len(values)!=count or
            any(type(value) is not int or value<=0 for value in values) or len(set(values))!=count):
        raise ListingStateError('Link IDs must be distinct positive integers, one per SPU title')
    return list(values)


def _read(path):
    p=Path(path);before=p.stat();data=p.read_bytes();after=p.stat()
    if (before.st_size,before.st_mtime_ns,before.st_ino)!=(after.st_size,after.st_mtime_ns,after.st_ino):
        raise ListingStateError(f'File changed during read: {p}')
    return data,hashlib.sha256(data).hexdigest()


def _path(value):
    if not isinstance(value,(str,os.PathLike)) or not str(value):
        raise ListingStateError('Expected an explicit filesystem path')
    return Path(value).resolve()


def _json(data):
    def unique(pairs):
        result={}
        for key,value in pairs:
            if key in result:raise ListingStateError(f'Duplicate JSON key: {key}')
            result[key]=value
        return result
    try:return json.loads(data,object_pairs_hook=unique)
    except (UnicodeError,json.JSONDecodeError) as exc:raise ListingStateError(f'Invalid JSON: {exc}') from exc


def _workbook(path,store_uid,*,require_links=False):
    import openpyxl
    p=_path(path)
    if p.suffix.lower()!='.xlsx':raise ListingStateError(f'Expected XLSX workbook: {p}')
    data,digest=_read(p)
    try:book=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=False)
    except Exception as exc:raise ListingStateError(f'Cannot read workbook {p}: {exc}') from exc
    try:
        if 'SPU' not in book.sheetnames:raise ListingStateError(f'Unknown workbook schema, missing SPU: {p}')
        sheet=book['SPU'];iterator=sheet.iter_rows();headers=next(iterator,())
        names=[unicodedata.normalize('NFKC',str(c.value)).strip().casefold() if c.value is not None else '' for c in headers]
        wanted=['店铺uid*','标题*']
        if any(names.count(key)!=1 for key in wanted):
            raise ListingStateError(f'Unknown or ambiguous SPU headers in {p}; need 店铺Uid* and 标题*')
        if require_links and names.count('链接序号')!=1:
            raise ListingStateError(f'Delivery identity requires one 链接序号 column: {p}')
        link_col=names.index('链接序号') if names.count('链接序号')==1 else None
        uid_col,title_col=[names.index(key) for key in wanted];records=[];stores=set();seen_ids=set()
        for row_no,cells in enumerate(iterator,2):
            if not any(c.value is not None for c in cells):continue
            if max(uid_col,title_col)>=len(cells):raise ListingStateError(f'Missing required SPU cells: {p}:{row_no}')
            if cells[uid_col].data_type=='f' or cells[title_col].data_type=='f':
                raise ListingStateError(f'Formula UID/title is not accepted: {p}:{row_no}')
            uid=cells[uid_col].value
            _uid(uid);title=cells[title_col].value
            normalized=_new_title(title) if require_links else normalize_title(title);stores.add(uid)
            link_cell=cells[link_col] if link_col is not None and link_col<len(cells) else None
            link_id=link_cell.value if link_cell is not None else None
            if require_links:
                if link_cell is None or link_cell.data_type=='f':raise ListingStateError(f'Unknown link identity: {p}:{row_no}')
                _link_ids([link_id],1)
                if link_id in seen_ids:raise ListingStateError(f'Duplicate link ID: {p}:{row_no}')
                seen_ids.add(link_id)
            if uid==store_uid:records.append({'title':title,'normalized':normalized,'path':str(p),'row':row_no,'sha256':digest,'link_id':link_id})
        return {'path':str(p),'sha256':digest,'records':records,'store_uids':sorted(stores)}
    finally:book.close()


def scan_titles(paths,store_uid,*,directory=None,pattern='*.xlsx'):
    """Scan explicit XLSX paths plus optional single-directory glob, read-only."""
    _uid(store_uid)
    if not isinstance(paths,(list,tuple)):raise ListingStateError('workbooks must be an explicit list')
    if not isinstance(pattern,str) or not pattern or '/' in pattern or '\\' in pattern or '**' in pattern:
        raise ListingStateError('History glob must stay in one directory')
    selected=list(paths)
    if directory is not None:
        folder=_path(directory)
        if not folder.is_dir():raise ListingStateError(f'History directory is inaccessible: {folder}')
        selected.extend(sorted(p for p in folder.glob(pattern) if p.is_file()))
    files=[];records=[];skipped=[];seen=set()
    for value in selected:
        p=_path(value)
        if p.name.startswith('~$'):skipped.append(str(p));continue
        key=os.path.normcase(str(p))
        if key in seen:continue
        seen.add(key);item=_workbook(p,store_uid);records.extend(item.pop('records'));files.append(item)
    return {'store_uid':store_uid,'titles':sorted({r['normalized'] for r in records}),
            'records':records,'files':files,'skipped_lock_files':skipped,
            'scope':'Only explicitly supplied files/current-directory glob; unknown history is not checked'}


def _empty():return {'schema_version':1,'revision':0,'stores':{}}


def _validate_state(state):
    if (not isinstance(state,dict) or set(state)!={'schema_version','revision','stores'}
            or type(state['schema_version']) is not int or state['schema_version']!=1
            or type(state['revision']) is not int or state['revision']<0 or not isinstance(state['stores'],dict)):
        raise ListingStateError('Unknown listing state schema')
    for uid,store in state['stores'].items():
        _uid(uid)
        if not isinstance(store,dict) or set(store)!={'titles','batches','latest'}:
            raise ListingStateError('Unknown store ledger schema')
        if not isinstance(store['titles'],dict) or not isinstance(store['batches'],dict) or not store['batches']:
            raise ListingStateError('Invalid store title/batch ledger')
        for batch,versions in store['batches'].items():
            _batch(batch)
            if not isinstance(versions,list) or not versions:raise ListingStateError('Empty batch version chain')
            previous=None
            for entry in versions:
                required={'path','sha256','store_uid','batch_id','parent_sha256','verification','titles','link_ids'}
                if not isinstance(entry,dict) or set(entry)!=required:raise ListingStateError('Unknown workbook pointer schema')
                _hash(entry['sha256']);_path(entry['path'])
                if entry['store_uid']!=uid or entry['batch_id']!=batch:raise ListingStateError('Pointer UID/batch mismatch')
                if entry['parent_sha256'] is not None:_hash(entry['parent_sha256'])
                if previous is not None and entry['parent_sha256']!=previous:raise ListingStateError('Broken batch parent chain')
                previous=entry['sha256'];verification=entry['verification']
                if not isinstance(verification,dict) or set(verification)!={'path','sha256'}:
                    raise ListingStateError('Unknown verification pointer schema')
                _path(verification['path']);_hash(verification['sha256'])
                titles=entry['titles']
                if (not isinstance(titles,list) or not titles or
                        any(not isinstance(t,str) for t in titles) or len(set(titles))!=len(titles)):
                    raise ListingStateError('Invalid pointer title set')
                if any(normalize_title(t)!=t or store['titles'].get(t)!=batch for t in titles):
                    raise ListingStateError('Pointer titles disagree with the batch title ledger')
                _link_ids(entry['link_ids'],len(titles))
        for title,batch in store['titles'].items():
            if normalize_title(title)!=title or not isinstance(batch,str) or batch not in store['batches']:
                raise ListingStateError('Invalid title owner')
        latest=store['latest']
        if not isinstance(latest,dict) or latest not in [v[-1] for v in store['batches'].values()]:
            raise ListingStateError('Latest pointer is stale or not a batch tip')
    return state


def _load(state_path,missing_ok=False):
    p=_path(state_path)
    if not p.exists():
        if missing_ok:return _empty(),None
        raise ListingStateError(f'Listing state is missing: {p}')
    raw,digest=_read(p);return _validate_state(_json(raw)),digest


def _verification(path,workbook_path,actual_sha256):
    raw,digest=_read(path);report=_json(raw)
    if (not isinstance(report,dict) or report.get('artifact_verified') is not True
            or report.get('artifact_sha256')!=actual_sha256 or report.get('sha256')!=actual_sha256
            or report.get('preserved_other_zip_parts') is not True or report.get('uploaded') is not False
            or report.get('operation') not in ('patch','batch') or type(report.get('changed_cells')) is not int
            or report['changed_cells']<0 or _path(report.get('output'))!=_path(workbook_path)):
        raise ListingStateError('Final Artifact verification does not bind this workbook/hash')
    return digest


def _resolve_entry(entry,store_uid):
    book=_workbook(entry['path'],store_uid,require_links=True)
    if book['sha256']!=entry['sha256']:raise ListingStateError('Tracked latest workbook drifted')
    if book['store_uids']!=[store_uid]:raise ListingStateError('Tracked workbook UID mismatch')
    if [r['normalized'] for r in book['records']]!=entry['titles']:
        raise ListingStateError('Tracked workbook titles disagree with state')
    if [r['link_id'] for r in book['records']]!=entry['link_ids']:
        raise ListingStateError('Tracked workbook link IDs disagree with state')
    if [r['row'] for r in book['records']]!=list(range(2,len(book['records'])+2)):
        raise ListingStateError('Tracked batch requires contiguous SPU row identity')
    verification=entry['verification']
    actual=_verification(verification['path'],entry['path'],book['sha256'])
    if actual!=verification['sha256']:raise ListingStateError('Tracked verification report drifted')
    return deepcopy(entry)


def resolve_latest(state_path,store_uid,batch_id=None,*,expected_sha256=None):
    """Resolve an explicit pointer by actual workbook and verification hashes."""
    _uid(store_uid);state,_=_load(state_path)
    store=state['stores'].get(store_uid)
    if store is None:raise ListingStateError('Unknown store in listing state')
    if batch_id is None:entry=store['latest']
    else:
        _batch(batch_id)
        if batch_id not in store['batches']:raise ListingStateError('Unknown batch in listing state')
        entry=store['batches'][batch_id][-1]
    if expected_sha256 is not None and _hash(expected_sha256)!=entry['sha256']:
        raise ListingStateError('Requested parent/latest version is stale')
    return _resolve_entry(entry,store_uid)


def _title_context(state,uid,batch_id,parent_sha256):
    if batch_id is not None:_batch(batch_id)
    if parent_sha256 is not None:_hash(parent_sha256)
    store=state['stores'].get(uid);versions=store['batches'].get(batch_id) if store else None
    if versions:
        if parent_sha256!=versions[-1]['sha256']:raise ListingStateError('Existing batch requires its current parent SHA')
        current=_resolve_entry(versions[-1],uid)
        return store,{v['sha256'] for v in versions},current['titles'],current['link_ids']
    if parent_sha256 is not None:
        if not store or parent_sha256!=store['latest']['sha256']:raise ListingStateError('Unknown or stale parent SHA')
        _resolve_entry(store['latest'],uid)
    return store,set(),[],[]


def _check(titles,uid,state,history,batch_id,parent_sha256,link_ids=None):
    if not isinstance(titles,(list,tuple)) or not titles:raise ListingStateError('Candidate titles must be a nonempty list')
    canonical=[_new_title(t) for t in titles]
    if len(set(canonical))!=len(canonical):raise ListingStateError('Candidate titles collide after NFKC normalization')
    if link_ids is not None:link_ids=_link_ids(link_ids,len(canonical))
    store,own_hashes,parent_titles,parent_ids=_title_context(state,uid,batch_id,parent_sha256);conflicts=[]
    if link_ids is None and any(i<len(parent_titles) and parent_titles[i]==t for i,t in enumerate(canonical)):
        raise ListingStateError('Current-title carryover requires confirmed link_ids; row-only identity is not accepted')
    carry_rows={title:index+2 for index,title in enumerate(canonical)
                if link_ids is not None and index<len(parent_titles) and parent_titles[index]==title
                and parent_ids[index]==link_ids[index]}
    for title in canonical:
        if store and title in store['titles']:
            owner=store['titles'][title]
            if owner!=batch_id or title not in carry_rows:
                conflicts.append({'title':title,'batch_id':owner,'source':'ledger'})
    for row in history['records']:
        if row['normalized'] in canonical:
            unchanged_own_row=(row['sha256'] in own_hashes and
                               carry_rows.get(row['normalized'])==row['row'] and link_ids is not None and
                               row['link_id']==link_ids[carry_rows[row['normalized']]-2])
            if not unchanged_own_row:conflicts.append(row)
    if conflicts:raise ListingStateError('Same-store title conflict: '+json.dumps(conflicts,ensure_ascii=False))
    return {'checked':True,'store_uid':uid,'candidate_count':len(canonical),'normalized_titles':canonical,
            'history_files':history['files'],'ledger_used':bool(store),'batch_id':batch_id,
            'unchanged_parent_title_rows':sorted(carry_rows.values()),
            'identity_binding':'Confirmed positive-integer link ID plus same contiguous SPU row/title',
            'link_ids_checked':link_ids is not None,
            'history_complete':False,'scope':history['scope']}


def check_titles(titles,store_uid,*,state_path=None,history_paths=(),batch_id=None,parent_sha256=None,link_ids=None):
    """Check candidate titles, permitting only hash-bound versions of this batch."""
    _uid(store_uid);history=scan_titles(history_paths,store_uid)
    state=_load(state_path,missing_ok=True)[0] if state_path is not None else _empty()
    return _check(titles,store_uid,state,history,batch_id,parent_sha256,link_ids)


def validate_title_history(titles,store_uid,history_config):
    """Config API for explicit workbooks/state or a nonrecursive directory glob."""
    allowed={'workbooks','directory','pattern','state_path','batch_id','parent_sha256','link_ids'}
    if not isinstance(history_config,dict) or set(history_config)-allowed:
        raise ListingStateError('Unknown title history configuration')
    history=scan_titles(history_config.get('workbooks',[]),store_uid,
                        directory=history_config.get('directory'),pattern=history_config.get('pattern','*.xlsx'))
    state=_load(history_config['state_path'],missing_ok=True)[0] if history_config.get('state_path') is not None else _empty()
    return _check(titles,_uid(store_uid),state,history,history_config.get('batch_id'),history_config.get('parent_sha256'),history_config.get('link_ids'))


@contextmanager
def _lock(path):
    lock=Path(str(path)+'.lock');token=(str(uuid.uuid4())+'\n').encode();fd=None;owner=None
    try:
        try:fd=os.open(lock,os.O_CREAT|os.O_EXCL|os.O_WRONLY|getattr(os,'O_BINARY',0),0o600)
        except FileExistsError as exc:raise ListingStateError(f'State is locked: {lock}') from exc
        info=os.fstat(fd);owner=(info.st_dev,info.st_ino)
        if os.write(fd,token)!=len(token):raise ListingStateError('Incomplete state lock write')
        os.fsync(fd);os.close(fd);fd=None
        yield
    finally:
        if fd is not None:os.close(fd)
        if owner is not None and lock.exists():
            current=lock.stat()
            if (current.st_dev,current.st_ino)==owner:lock.unlink()


def _atomic_state(path,state,previous_sha):
    payload=json.dumps(_validate_state(state),ensure_ascii=False,sort_keys=True,indent=2).encode('utf-8')
    fd,name=tempfile.mkstemp(prefix=path.name+'.',suffix='.tmp',dir=path.parent)
    try:
        with os.fdopen(fd,'wb') as handle:handle.write(payload);handle.flush();os.fsync(handle.fileno())
        current=_read(path)[1] if path.exists() else None
        if current!=previous_sha:raise ListingStateError('State changed despite lock; refusing overwrite')
        os.replace(name,path)
    finally:
        if os.path.exists(name):os.unlink(name)


def register_delivery(state_path,workbook_path,verification_path,store_uid,batch_id,*,parent_sha256=None,history_paths=()):
    """Register only a final Artifact-verified delivery; commit ledger+latest once."""
    _uid(store_uid);_batch(batch_id);path=_path(state_path)
    module_root=Path(__file__).resolve().parents[1]
    if path==module_root or module_root in path.parents:
        raise ListingStateError('State must be stored in a batch/ledger directory outside the skill tree')
    book=_workbook(workbook_path,store_uid,require_links=True)
    if book['store_uids']!=[store_uid] or not book['records']:raise ListingStateError('Delivery workbook must contain only the confirmed store UID')
    if [r['row'] for r in book['records']]!=list(range(2,len(book['records'])+2)):
        raise ListingStateError('Delivery requires contiguous SPU row identity; adapt link IDs explicitly before reordering')
    verification_path=_path(verification_path);verification_sha=_verification(verification_path,book['path'],book['sha256'])
    titles=[r['normalized'] for r in book['records']];history=scan_titles(history_paths,store_uid)
    path.parent.mkdir(parents=True,exist_ok=True)
    with _lock(path):
        actual_ids=[r['link_id'] for r in book['records']]
        state,old_sha=_load(path,missing_ok=True);audit=_check(titles,store_uid,state,history,batch_id,parent_sha256,actual_ids)
        if _read(book['path'])[1]!=book['sha256'] or _read(verification_path)[1]!=verification_sha:
            raise ListingStateError('Output or verification changed before state commit')
        updated=deepcopy(state);updated['revision']+=1
        store=updated['stores'].setdefault(store_uid,{'titles':{},'batches':{},'latest':None})
        entry={'path':book['path'],'sha256':book['sha256'],'store_uid':store_uid,'batch_id':batch_id,
               'parent_sha256':parent_sha256,'verification':{'path':str(verification_path),'sha256':verification_sha},'titles':titles,'link_ids':actual_ids}
        for title in titles:store['titles'][title]=batch_id
        store['batches'].setdefault(batch_id,[]).append(entry);store['latest']=deepcopy(entry)
        _atomic_state(path,updated,old_sha)
    return {'state_path':str(path),'revision':updated['revision'],'latest':entry,'title_audit':audit,'registered':True}


def main(argv=None):
    parser=argparse.ArgumentParser(description=__doc__);sub=parser.add_subparsers(dest='command',required=True)
    for name in ('scan','check','resolve','register'):
        p=sub.add_parser(name);p.add_argument('--store-uid',required=True)
        if name!='scan':p.add_argument('--state',required=name in ('resolve','register'))
        if name in ('check','resolve','register'):p.add_argument('--batch-id',required=name=='register')
        if name in ('check','register'):p.add_argument('--parent-sha256')
        if name in ('scan','check'):
            p.add_argument('--workbook',action='append',default=[]);p.add_argument('--directory');p.add_argument('--pattern',default='*.xlsx')
        if name=='check':
            p.add_argument('--titles',required=True,help='JSON array of candidate titles')
            p.add_argument('--link-ids',help='JSON array of candidate positive-integer link IDs')
        if name=='resolve':p.add_argument('--expected-sha256')
        if name=='register':
            p.add_argument('--workbook',required=True);p.add_argument('--verification',required=True)
            p.add_argument('--history-workbook',action='append',default=[])
    args=parser.parse_args(argv)
    try:
        if args.command=='scan':result=scan_titles(args.workbook,args.store_uid,directory=args.directory,pattern=args.pattern)
        elif args.command=='resolve':result=resolve_latest(args.state,args.store_uid,args.batch_id,expected_sha256=args.expected_sha256)
        elif args.command=='check':
            titles=_json(Path(args.titles).read_bytes())
            result=validate_title_history(titles,args.store_uid,{'workbooks':args.workbook,'directory':args.directory,'pattern':args.pattern,
                'state_path':args.state,'batch_id':args.batch_id,'parent_sha256':args.parent_sha256,
                'link_ids':_json(Path(args.link_ids).read_bytes()) if args.link_ids else None})
        else:result=register_delivery(args.state,args.workbook,args.verification,args.store_uid,args.batch_id,
                                     parent_sha256=args.parent_sha256,history_paths=args.history_workbook)
        print(json.dumps(result,ensure_ascii=False));return 0
    except (ValueError,OSError) as exc:
        print(json.dumps({'error':str(exc)},ensure_ascii=False),file=sys.stderr);return 1


if __name__=='__main__':
    if hasattr(sys.stdout,'reconfigure'):sys.stdout.reconfigure(encoding='utf-8');sys.stderr.reconfigure(encoding='utf-8')
    raise SystemExit(main())
